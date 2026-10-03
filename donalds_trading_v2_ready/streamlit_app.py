from pathlib import Path
from io import BytesIO
import re
import math
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Donald's Trading — Performance Lab", page_icon="📈", layout="wide", initial_sidebar_state="expanded")

PURPLE = "#7C3AED"
PURPLE_DARK = "#5B21B6"
PURPLE_SOFT = "#EDE9FE"
INK = "#17131F"
MUTED = "#6B6474"
GREEN = "#0F9D6E"
RED = "#D6455D"
GRID = "#ECE8F1"
APP_DIR = Path(__file__).resolve().parent
DEFAULT_FILE = APP_DIR / "trading_data.xlsx"

st.markdown("""
<style>
#MainMenu, footer {visibility:hidden}
header[data-testid="stHeader"] {background:rgba(248,247,252,.92); backdrop-filter:blur(10px)}
.block-container {padding-top:1.6rem; padding-bottom:3rem; max-width:1500px}
[data-testid="stSidebar"] {background:#FFFFFF; border-right:1px solid #ECE8F1}
[data-testid="stSidebar"] .block-container {padding-top:1.4rem}
h1,h2,h3 {letter-spacing:-.035em; color:#17131F}
p,div,span,label {color:#332D3B}
.hero {background:linear-gradient(120deg,#FFFFFF 0%,#FAF8FF 58%,#F0E9FF 100%);border:1px solid #E7DDFB;border-radius:28px;padding:30px 34px;margin-bottom:20px;box-shadow:0 12px 35px rgba(70,40,120,.07)}
.eyebrow {font-size:.75rem;font-weight:800;letter-spacing:.12em;text-transform:uppercase;color:#7C3AED;margin-bottom:8px}
.hero h1 {font-size:2.5rem;line-height:1.05;margin:0 0 10px}.hero p{font-size:1.05rem;color:#6B6474;max-width:850px;margin:0}
.kpi {background:#FFFFFF;border:1px solid #EDE9F2;border-radius:20px;padding:18px 19px;min-height:132px;box-shadow:0 8px 24px rgba(43,31,60,.045)}
.kpi .label{font-size:.78rem;font-weight:700;color:#7B7383;text-transform:uppercase;letter-spacing:.05em}.kpi .value{font-size:1.75rem;font-weight:800;color:#17131F;margin:7px 0 3px}.kpi .sub{font-size:.82rem;color:#7B7383}.purple-pill{display:inline-block;background:#EDE9FE;color:#6D28D9;border-radius:999px;padding:5px 10px;font-weight:800;font-size:.75rem}
.section-card {background:#FFFFFF;border:1px solid #EDE9F2;border-radius:22px;padding:20px;box-shadow:0 8px 24px rgba(43,31,60,.035);margin-bottom:14px}
.insight {background:#F5F1FF;border-left:4px solid #7C3AED;border-radius:12px;padding:13px 15px;margin:8px 0;color:#3D3151}
.small-note {font-size:.82rem;color:#7B7383}
div[data-testid="stMetric"] {background:#fff;border:1px solid #EDE9F2;padding:14px 16px;border-radius:16px}
.stTabs [data-baseweb="tab-list"] {gap:8px}.stTabs [data-baseweb="tab"]{background:#fff;border:1px solid #E9E4EF;border-radius:12px;padding:8px 15px}.stTabs [aria-selected="true"]{background:#EDE9FE!important;color:#5B21B6!important}
.stButton>button,.stDownloadButton>button {border-radius:12px;border:1px solid #D8CCF3;font-weight:700}
.stButton>button:hover,.stDownloadButton>button:hover {border-color:#7C3AED;color:#6D28D9}
[data-testid="stDataFrame"] {border:1px solid #EDE9F2;border-radius:16px;overflow:hidden}
</style>
""", unsafe_allow_html=True)

DAY_WORDS = ["Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"]

def money(v):
    if v is None or pd.isna(v): return "—"
    return f"R {float(v):,.0f}"

def pct(v, digits=1):
    if v is None or pd.isna(v): return "—"
    return f"{float(v)*100:.{digits}f}%"

def n(v):
    try:
        if v in (None, "", "-"): return 0.0
        return float(v)
    except Exception: return 0.0

def clean_text(v): return "" if v is None else str(v).strip()

def sheet_week_label(name):
    return name.replace(". 2026", " 2026").replace(". 26", " 2026")

def extract_day(text):
    t = clean_text(text)
    for d in DAY_WORDS:
        if t.lower().startswith(d.lower()): return d
    return None

def header_map(row):
    out={}
    for i,v in enumerate(row):
        s=clean_text(v).lower().replace("  "," ")
        if s in ("trade","trades"): out["trade"]=i
        elif s=="coin": out["coin"]=i
        elif "capital allocated" in s: out["capital"]=i
        elif "target stop loss" in s: out["stop"]=i
        elif "profit target" in s: out["target"]=i
        elif "actual profit (%)" in s: out["profit_pct"]=i
        elif "actual profit (zar)" in s: out["profit_zar"]=i
        elif "actual loss (%)" in s: out["loss_pct"]=i
        elif "actual loss (zar)" in s: out["loss_zar"]=i
        elif "total (zar)" in s: out["total"]=i
    return out

def parse_summary(ws, week_name):
    rows=list(ws.iter_rows(values_only=True))
    if not rows or clean_text(rows[0][0]).lower()!="metric": return []
    header=list(rows[0])
    result=[]
    for r in rows[1:20]:
        metric=clean_text(r[0])
        if not metric: continue
        for j in range(1,min(len(header),len(r))):
            h=header[j]
            if h is None: continue
            label=h.strftime("%a %d %b") if hasattr(h,"strftime") else clean_text(h)
            if label.upper() in ("TOTAL","WEEKLY RESULTS") or not label: continue
            result.append({"Week":week_name,"Day":label,"Metric":metric,"Value":r[j]})
    return result

def parse_workbook(source):
    wb=load_workbook(source, data_only=True, read_only=True)
    trades=[]; summaries=[]; audit=[]; model={}
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        if ws.title.lower()=="profitability model":
            for r in rows[:18]:
                key=clean_text(r[0])
                if key: model[key]=list(r[1:])
            continue
        week=sheet_week_label(ws.title)
        summaries.extend(parse_summary(ws,week))
        current_day=None; hmap=None; trade_count=0
        for r_idx,row in enumerate(rows, start=1):
            first=clean_text(row[0] if len(row)>0 else None)
            maybe_day=extract_day(first)
            if maybe_day and not first.lower().startswith(("trade","trades")):
                current_day=maybe_day
            hm=header_map(row)
            if "trade" in hm and "coin" in hm and "capital" in hm:
                hmap=hm; continue
            if hmap:
                ti=hmap.get("trade",0)
                tv=row[ti] if ti<len(row) else None
                try: trade_no=int(float(tv))
                except Exception:
                    if first.upper()=="TOTAL" or (first and not maybe_day): hmap=None
                    continue
                def val(key):
                    idx=hmap.get(key)
                    return row[idx] if idx is not None and idx<len(row) else None
                coin=clean_text(val("coin"))
                if not coin: continue
                profit_z=n(val("profit_zar")); loss_z=n(val("loss_zar"))
                # Some later sheets put actual profit ZAR farther right; derive from percentages if missing.
                capital=n(val("capital")); profit_p=n(val("profit_pct")); loss_p=n(val("loss_pct"))
                if profit_z==0 and profit_p!=0: profit_z=capital*profit_p
                if loss_z==0 and loss_p!=0: loss_z=capital*loss_p
                pnl=profit_z+loss_z
                target_pct=n(val("target")); stop_pct=n(val("stop"))
                trades.append({"Week":week,"Day":current_day or "Unspecified","Trade":trade_no,"Coin":coin,
                    "Capital":capital,"Stop Loss %":stop_pct,"Target %":target_pct,"Profit %":profit_p,
                    "Profit ZAR":profit_z,"Loss %":loss_p,"Loss ZAR":loss_z,"P&L":pnl,"Source Row":r_idx})
                trade_count+=1
        audit.append({"Sheet":week,"Trades parsed":trade_count,"Summary metrics":sum(1 for x in summaries if x["Week"]==week)})
    t=pd.DataFrame(trades)
    s=pd.DataFrame(summaries)
    return t,s,model,pd.DataFrame(audit),wb.sheetnames

@st.cache_data(show_spinner=False)
def load_default():
    if not DEFAULT_FILE.exists(): return None
    return parse_workbook(DEFAULT_FILE)

@st.cache_data(show_spinner=False)
def load_uploaded(raw): return parse_workbook(BytesIO(raw))

def plotly_layout(fig, height=390):
    fig.update_layout(height=height,margin=dict(l=15,r=15,t=45,b=15),paper_bgcolor="white",plot_bgcolor="white",font=dict(color=INK),
                      title_font=dict(size=18),legend_title_text="",hoverlabel=dict(bgcolor="white"))
    fig.update_xaxes(showgrid=False,linecolor=GRID)
    fig.update_yaxes(gridcolor=GRID,zerolinecolor=GRID)
    return fig

def kpi(label,value,sub=""):
    st.markdown(f'<div class="kpi"><div class="label">{label}</div><div class="value">{value}</div><div class="sub">{sub}</div></div>',unsafe_allow_html=True)

def metric_lookup(summary, contains):
    if summary.empty: return pd.DataFrame()
    m=summary[summary["Metric"].str.lower().str.contains(contains.lower(),regex=False,na=False)].copy()
    return m

def model_first(model,key,default=0):
    vals=model.get(key,[])
    for v in vals:
        if isinstance(v,(int,float)) and not isinstance(v,bool): return float(v)
    return default

# Sidebar / source
with st.sidebar:
    st.markdown("## 📈 Performance Lab")
    st.caption("Donald's trading journal, translated into a clear performance story.")
    st.markdown("---")
    uploaded=st.file_uploader("Use a different workbook",type=["xlsx"],help="Optional. The supplied workbook is already bundled with the app.")
    page=st.radio("Explore",["Overview","Progress & Weeks","Trade Journal","Coins","Risk & Targets","Profitability Simulator","Report & Project"],label_visibility="collapsed")
    st.markdown("---")
    st.caption("Light portfolio build • Streamlit + Python + Plotly + Excel")

try:
    parsed=load_uploaded(uploaded.getvalue()) if uploaded else load_default()
except Exception as e:
    st.error("I couldn't read that workbook. Please use the original .xlsx trading sheet.")
    st.code(str(e)); st.stop()
if parsed is None:
    st.error("Bundled workbook not found. Keep trading_data.xlsx beside streamlit_app.py."); st.stop()
trades,summary,model,audit,sheet_names=parsed
if trades.empty:
    st.error("No trade rows were detected in the workbook."); st.stop()

# Canonical metrics
trades["Result"]=trades["P&L"].apply(lambda x:"Win" if x>0 else ("Loss" if x<0 else "Flat"))
trades["Target Hit"]=trades.apply(lambda r: r["Profit %"]>=r["Target %"] and r["Target %"]>0,axis=1)
trades["R Multiple"]=trades.apply(lambda r: (r["P&L"]/(r["Capital"]*r["Stop Loss %"])) if r["Capital"]>0 and r["Stop Loss %"]>0 else 0,axis=1)
week_order=list(dict.fromkeys(trades["Week"].tolist()))
trades["WeekOrder"]=trades["Week"].map({w:i for i,w in enumerate(week_order)})
trades=trades.sort_values(["WeekOrder","Day","Trade"]).reset_index(drop=True)

start_cap=model_first(model,"My trading capital (TC)  (ZAR)",0)
balance_rows=metric_lookup(summary,"actual balance")
current_balance=0
if not balance_rows.empty:
    vals=pd.to_numeric(balance_rows["Value"],errors="coerce").dropna()
    if len(vals): current_balance=float(vals.iloc[-1])
if start_cap==0:
    tc=metric_lookup(summary,"trading capital")
    vals=pd.to_numeric(tc["Value"],errors="coerce").dropna() if not tc.empty else pd.Series(dtype=float)
    start_cap=float(vals.iloc[0]) if len(vals) else float(trades["Capital"].max())
total_profit=float(trades["P&L"].sum())
if current_balance==0: current_balance=start_cap+total_profit
wins=int((trades["P&L"]>0).sum()); losses=int((trades["P&L"]<0).sum()); executed=wins+losses
win_rate=wins/executed if executed else 0
coin_perf=trades.groupby("Coin",as_index=False).agg(PnL=("P&L","sum"),Trades=("Trade","count"),Wins=("Result",lambda x:(x=="Win").sum()))
coin_perf["Win rate"]=coin_perf["Wins"]/coin_perf["Trades"]
best_coin=coin_perf.sort_values("PnL",ascending=False).iloc[0]
target_hits=int(trades["Target Hit"].sum()); target_rate=target_hits/len(trades) if len(trades) else 0

# Page hero
st.markdown(f'''<div class="hero"><div class="eyebrow">Trading performance portfolio</div><h1>{page}</h1><p>From spreadsheet rows to a clear story: performance, targets, risk, coins and progression — using the data already recorded in your workbook.</p></div>''',unsafe_allow_html=True)

if page=="Overview":
    c1,c2,c3,c4,c5,c6=st.columns(6)
    with c1:kpi("Starting capital",money(start_cap),"Model trading capital")
    with c2:kpi("Current balance",money(current_balance),"Latest recorded balance")
    with c3:kpi("Total trade P&L",money(total_profit),f"Across {len(trades):,} parsed trades")
    with c4:kpi("Win rate",pct(win_rate),f"{wins} wins • {losses} losses")
    with c5:kpi("Best coin",best_coin["Coin"],f"{money(best_coin['PnL'])} total P&L")
    with c6:kpi("Targets hit",pct(target_rate),f"{target_hits} trades reached target")
    st.write("")
    left,right=st.columns([1.55,1])
    with left:
        weekly=trades.groupby("Week",sort=False,as_index=False).agg(PnL=("P&L","sum"),Trades=("Trade","count"))
        weekly["Cumulative P&L"]=weekly["PnL"].cumsum(); weekly["Projected balance"]=start_cap+weekly["Cumulative P&L"]
        fig=go.Figure()
        fig.add_trace(go.Scatter(x=weekly["Week"],y=weekly["Projected balance"],mode="lines+markers",name="Balance path",line=dict(color=PURPLE,width=4),fill="tozeroy",fillcolor="rgba(124,58,237,.08)"))
        fig.update_layout(title="Growth curve from recorded trades",yaxis_title="ZAR")
        st.plotly_chart(plotly_layout(fig,410),use_container_width=True)
    with right:
        res=trades[trades["Result"]!="Flat"]["Result"].value_counts().rename_axis("Result").reset_index(name="Trades")
        fig=px.pie(res,names="Result",values="Trades",hole=.68,color="Result",color_discrete_map={"Win":GREEN,"Loss":RED})
        fig.update_layout(title="Wins vs losses",annotations=[dict(text=f"{win_rate:.0%}<br><span style='font-size:11px'>win rate</span>",x=.5,y=.5,showarrow=False,font=dict(size=22))])
        st.plotly_chart(plotly_layout(fig,410),use_container_width=True)
    a,b=st.columns(2)
    with a:
        day=trades.groupby(["Week","Day"],sort=False,as_index=False)["P&L"].sum()
        fig=px.bar(day,x="Day",y="P&L",color="Week",barmode="group",title="P&L by trading day",color_discrete_sequence=[PURPLE,"#A78BFA","#C4B5FD","#6D28D9","#8B5CF6","#DDD6FE","#4C1D95"])
        st.plotly_chart(plotly_layout(fig),use_container_width=True)
    with b:
        top=coin_perf.sort_values("PnL",ascending=False).head(10).sort_values("PnL")
        fig=px.bar(top,x="PnL",y="Coin",orientation="h",title="Top coins by recorded P&L",color="PnL",color_continuous_scale=["#EDE9FE",PURPLE_DARK])
        fig.update_coloraxes(showscale=False); st.plotly_chart(plotly_layout(fig),use_container_width=True)
    st.markdown(f'<div class="insight"><b>Quick read:</b> The workbook contains <b>{len(trades):,} individual trades</b> across <b>{len(week_order)} recorded weeks</b>. Your strongest coin by total recorded P&L is <b>{best_coin["Coin"]}</b>, and <b>{target_rate:.0%}</b> of parsed trades met or exceeded their recorded profit target.</div>',unsafe_allow_html=True)

elif page=="Progress & Weeks":
    tab1,tab2,tab3=st.tabs(["Progress timeline","Week explorer","Target vs actual"])
    with tab1:
        weekly=trades.groupby("Week",sort=False,as_index=False).agg(PnL=("P&L","sum"),Trades=("Trade","count"),Wins=("Result",lambda x:(x=="Win").sum()),TargetHits=("Target Hit","sum"))
        weekly["Win rate"]=weekly["Wins"]/weekly["Trades"]; weekly["Cumulative P&L"]=weekly["PnL"].cumsum()
        fig=px.bar(weekly,x="Week",y="PnL",text_auto=".0f",title="Weekly P&L progression",color="PnL",color_continuous_scale=["#F2DCE3","#F8F7FC",PURPLE])
        fig.update_coloraxes(showscale=False); st.plotly_chart(plotly_layout(fig,430),use_container_width=True)
        st.dataframe(weekly[["Week","Trades","PnL","Win rate","TargetHits","Cumulative P&L"]].style.format({"PnL":"R {:,.0f}","Win rate":"{:.1%}","Cumulative P&L":"R {:,.0f}"}),use_container_width=True,hide_index=True)
    with tab2:
        chosen=st.selectbox("Choose a week",week_order,index=len(week_order)-1)
        w=trades[trades["Week"]==chosen]
        wwin=(w["P&L"]>0).mean(); wpnl=w["P&L"].sum(); whits=w["Target Hit"].mean()
        c1,c2,c3,c4=st.columns(4)
        c1.metric("Trades",len(w)); c2.metric("Week P&L",money(wpnl)); c3.metric("Win rate",pct(wwin)); c4.metric("Target hit rate",pct(whits))
        day=w.groupby("Day",sort=False,as_index=False).agg(PnL=("P&L","sum"),Trades=("Trade","count"))
        fig=px.bar(day,x="Day",y="PnL",text_auto=".0f",title=f"{chosen}: day-by-day result",color="PnL",color_continuous_scale=["#F6D7DF","#F8F7FC",PURPLE])
        fig.update_coloraxes(showscale=False); st.plotly_chart(plotly_layout(fig),use_container_width=True)
    with tab3:
        target=metric_lookup(summary,"target profit")
        actual=metric_lookup(summary,"profit after loss")
        if actual.empty: actual=metric_lookup(summary,"actual profit/loss (zar)")
        if not target.empty and not actual.empty:
            x=target[["Week","Day","Value"]].rename(columns={"Value":"Target"}).merge(actual[["Week","Day","Value"]].rename(columns={"Value":"Actual"}),on=["Week","Day"],how="inner")
            x["Target"]=pd.to_numeric(x["Target"],errors="coerce"); x["Actual"]=pd.to_numeric(x["Actual"],errors="coerce"); x=x.dropna()
            fig=go.Figure([go.Bar(name="Target",x=x["Week"]+" • "+x["Day"],y=x["Target"],marker_color="#D8CCF3"),go.Bar(name="Actual",x=x["Week"]+" • "+x["Day"],y=x["Actual"],marker_color=PURPLE)])
            fig.update_layout(barmode="group",title="Recorded target vs actual profit"); st.plotly_chart(plotly_layout(fig,470),use_container_width=True)
        else: st.info("The workbook does not contain matching summary target/actual rows for every week, so trade-level target attainment is shown on the Overview instead.")

elif page=="Trade Journal":
    f1,f2,f3=st.columns(3)
    weeks=f1.multiselect("Weeks",week_order,default=week_order)
    coins=f2.multiselect("Coins",sorted(trades["Coin"].unique()))
    results=f3.multiselect("Result",["Win","Loss","Flat"],default=["Win","Loss","Flat"])
    view=trades[trades["Week"].isin(weeks)&trades["Result"].isin(results)].copy()
    if coins:view=view[view["Coin"].isin(coins)]
    show=view[["Week","Day","Trade","Coin","Capital","Stop Loss %","Target %","Profit %","Profit ZAR","Loss %","Loss ZAR","P&L","Result","Target Hit","R Multiple"]]
    st.dataframe(show.style.format({"Capital":"R {:,.0f}","Stop Loss %":"{:.2%}","Target %":"{:.2%}","Profit %":"{:.2%}","Profit ZAR":"R {:,.2f}","Loss %":"{:.2%}","Loss ZAR":"R {:,.2f}","P&L":"R {:,.2f}","R Multiple":"{:.2f}R"}),use_container_width=True,hide_index=True,height=620)
    st.download_button("Download filtered journal (CSV)",show.to_csv(index=False).encode(),"donalds_trade_journal.csv","text/csv")

elif page=="Coins":
    cp=coin_perf.sort_values("PnL",ascending=False)
    c1,c2=st.columns([1.25,1])
    with c1:
        fig=px.scatter(cp,x="Trades",y="PnL",size="Trades",color="Win rate",hover_name="Coin",title="Coin performance map",color_continuous_scale=["#DDD6FE",PURPLE_DARK])
        st.plotly_chart(plotly_layout(fig,500),use_container_width=True)
    with c2:
        st.markdown("### Coin leaderboard")
        st.dataframe(cp.style.format({"PnL":"R {:,.0f}","Win rate":"{:.1%}"}),use_container_width=True,hide_index=True,height=470)
    selected=st.selectbox("Inspect one coin",cp["Coin"].tolist())
    ct=trades[trades["Coin"]==selected]
    c1,c2,c3,c4=st.columns(4); c1.metric("Trades",len(ct)); c2.metric("P&L",money(ct["P&L"].sum())); c3.metric("Win rate",pct((ct["P&L"]>0).mean())); c4.metric("Targets hit",pct(ct["Target Hit"].mean()))

elif page=="Risk & Targets":
    c1,c2,c3,c4=st.columns(4)
    avg_r=trades[trades["R Multiple"].replace([math.inf,-math.inf],pd.NA).notna()]["R Multiple"].mean()
    c1.metric("Average stop loss",pct(trades["Stop Loss %"].replace(0,pd.NA).mean(),2)); c2.metric("Average target",pct(trades["Target %"].replace(0,pd.NA).mean(),2)); c3.metric("Average R multiple",f"{avg_r:.2f}R"); c4.metric("Target hit rate",pct(target_rate))
    a,b=st.columns(2)
    with a:
        fig=px.histogram(trades,x="R Multiple",nbins=30,title="Distribution of trade R-multiples",color_discrete_sequence=[PURPLE]); st.plotly_chart(plotly_layout(fig),use_container_width=True)
    with b:
        alloc=trades.groupby("Week",sort=False,as_index=False)["Capital"].sum()
        fig=px.bar(alloc,x="Week",y="Capital",title="Capital allocated by recorded week",color_discrete_sequence=[PURPLE]); st.plotly_chart(plotly_layout(fig),use_container_width=True)
    best=trades.nlargest(8,"P&L")[["Week","Day","Coin","Capital","P&L","R Multiple"]]; worst=trades.nsmallest(8,"P&L")[["Week","Day","Coin","Capital","P&L","R Multiple"]]
    l,r=st.columns(2)
    with l: st.markdown("### Best trades"); st.dataframe(best.style.format({"Capital":"R {:,.0f}","P&L":"R {:,.2f}","R Multiple":"{:.2f}R"}),hide_index=True,use_container_width=True)
    with r: st.markdown("### Worst trades"); st.dataframe(worst.style.format({"Capital":"R {:,.0f}","P&L":"R {:,.2f}","R Multiple":"{:.2f}R"}),hide_index=True,use_container_width=True)

elif page=="Profitability Simulator":
    st.markdown("### Turn the workbook's profitability idea into a live model")
    st.caption("Change the assumptions below. The simulator shows a scenario — it does not predict future returns.")
    base_cap=start_cap if start_cap>0 else 300000
    a,b,c,d=st.columns(4)
    sim_cap=a.number_input("Starting capital (ZAR)",min_value=1000.0,value=float(base_cap),step=1000.0)
    allocation=b.number_input("Allocation per trade",min_value=100.0,value=float(model_first(model,"Allocation/ Trade (ZAR) ",30000)),step=500.0)
    trades_day=c.slider("Trades per day",1,50,int(model_first(model,"No of executed trades",10)))
    days=d.slider("Trading days",5,120,20)
    e,f,g=st.columns(3)
    sim_win=e.slider("Win rate",0.0,1.0,float(min(max(win_rate,.01),.99)),.01)
    sim_target=f.slider("Profit target per winning trade",0.001,0.10,float(max(trades["Target %"].replace(0,pd.NA).median(),.02)),.001,format="%.3f")
    sim_stop=g.slider("Stop loss per losing trade",0.001,0.05,float(max(trades["Stop Loss %"].replace(0,pd.NA).median(),.005)),.001,format="%.3f")
    rows=[]; bal=sim_cap
    for day in range(1,days+1):
        expected_wins=trades_day*sim_win; expected_losses=trades_day*(1-sim_win)
        daily=expected_wins*allocation*sim_target-expected_losses*allocation*sim_stop
        bal+=daily; rows.append({"Day":day,"Projected balance":bal,"Projected P&L":daily})
    sim=pd.DataFrame(rows); projected_profit=bal-sim_cap
    c1,c2,c3,c4=st.columns(4); c1.metric("Projected end balance",money(bal)); c2.metric("Projected profit",money(projected_profit)); c3.metric("Expected daily P&L",money(sim["Projected P&L"].iloc[0])); c4.metric("Scenario return",pct(projected_profit/sim_cap if sim_cap else 0))
    fig=px.line(sim,x="Day",y="Projected balance",title="Scenario growth path",color_discrete_sequence=[PURPLE]); fig.update_traces(line_width=4); st.plotly_chart(plotly_layout(fig,430),use_container_width=True)
    st.markdown('<div class="insight"><b>How to read it:</b> this converts your target/stop-loss framework into expected-value math using the assumptions you choose. It is a planning visual, not a guarantee of trading performance.</div>',unsafe_allow_html=True)

elif page=="Report & Project":
    tab1,tab2,tab3=st.tabs(["Executive report","Workbook audit","About the build"])
    with tab1:
        report=pd.DataFrame({"Metric":["Starting capital","Current balance","Recorded trade P&L","Trades parsed","Wins","Losses","Win rate","Best-performing coin","Best coin P&L","Trade targets hit"],"Value":[money(start_cap),money(current_balance),money(total_profit),str(len(trades)),str(wins),str(losses),pct(win_rate),best_coin["Coin"],money(best_coin["PnL"]),pct(target_rate)]})
        st.dataframe(report,use_container_width=True,hide_index=True)
        csv=report.to_csv(index=False).encode(); st.download_button("Download executive summary",csv,"donalds_trading_executive_summary.csv","text/csv")
        st.markdown("### What the data says")
        st.markdown(f"- **{len(trades):,} trades** were parsed from the weekly journal sheets.\n- **{wins}** were profitable and **{losses}** recorded losses, producing a **{win_rate:.1%} win rate**.\n- **{best_coin['Coin']}** has the highest aggregate recorded trade P&L at **{money(best_coin['PnL'])}**.\n- **{target_hits} trades ({target_rate:.1%})** met or exceeded their recorded profit target.")
    with tab2:
        st.dataframe(audit,use_container_width=True,hide_index=True)
        st.caption("Profitability model is read separately because it is a planning/model sheet rather than a historical week.")
        st.write("**Workbook sheets detected:**",", ".join(sheet_names))
    with tab3:
        st.markdown("""### Why this is a programming project
This app does more than display an Excel file. It **parses differently structured weekly sheets**, normalises individual trades, calculates performance metrics, creates interactive filters and charts, models risk/targets, and turns the profitability assumptions into a live simulator.

**Stack:** Python • Streamlit • Plotly • pandas • openpyxl • GitHub-ready project structure.

**Portfolio story:** *I took a real trading workbook and converted it into an interactive analytics product that makes the data easier to understand, explore and present.*
""")
